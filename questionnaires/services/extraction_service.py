# ============================================================================
# FILE: questionnaires/services/extraction_service.py
# ============================================================================

import json
import PyPDF2
import docx
import openpyxl
from typing import List, Dict
from django.conf import settings

class QuestionnaireExtractor:
    """Extract questions from uploaded files using Claude AI"""
    
    def __init__(self, api_key=None):
        self.api_key = api_key or getattr(settings, 'ANTHROPIC_API_KEY', None)
        if not self.api_key:
            raise ValueError("ANTHROPIC_API_KEY not found in settings")
        
        try:
            import anthropic
            self.client = anthropic.Anthropic(api_key=self.api_key)
        except ImportError:
            raise ImportError("Please install anthropic: pip install anthropic")
    
    def extract_text_from_file(self, file_path: str) -> str:
        """Extract text content from various file formats"""
        extension = file_path.lower().split('.')[-1]
        
        if extension == 'pdf':
            return self._extract_from_pdf(file_path)
        elif extension in ['docx', 'doc']:
            return self._extract_from_docx(file_path)
        elif extension in ['xlsx', 'xls']:
            return self._extract_from_excel(file_path)
        elif extension == 'txt':
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read()
        else:
            raise ValueError(f"Unsupported file format: {extension}")
    
    def _extract_from_pdf(self, file_path: str) -> str:
        """Extract text from PDF"""
        text = ""
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                for page in pdf_reader.pages:
                    text += page.extract_text() + "\n"
        except Exception as e:
            raise Exception(f"Error extracting PDF: {str(e)}")
        return text
    
    def _extract_from_docx(self, file_path: str) -> str:
        """Extract text from DOCX"""
        try:
            doc = docx.Document(file_path)
            return "\n".join([paragraph.text for paragraph in doc.paragraphs])
        except Exception as e:
            raise Exception(f"Error extracting DOCX: {str(e)}")
    
    def _extract_from_excel(self, file_path: str) -> str:
        """Extract text from Excel"""
        try:
            workbook = openpyxl.load_workbook(file_path)
            text = ""
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    text += " ".join([str(cell) for cell in row if cell]) + "\n"
            return text
        except Exception as e:
            raise Exception(f"Error extracting Excel: {str(e)}")
    
    def extract_questions_with_ai(self, content: str, question_types: List[str]) -> Dict:
        """Use Claude AI to extract and generate questions"""
        
        prompt = f"""You are an educational content analyzer. Extract and generate questions from the following educational material.

CONTENT:
{content[:8000]}

TASK:
Analyze this content and generate questions of the following types:
{', '.join(question_types)}

For each question type requested, create 5-10 high-quality questions based on the content.

QUESTION TYPES GUIDE:
- multiple_choice: Provide question, 4 options (A-D), and correct answer
- true_false: Provide statement and correct answer (true/false)
- identification: Provide question requiring a specific term/concept
- essay: Provide open-ended analytical questions
- fill_blank: Provide sentences with blanks and answers
- matching: Provide two columns to match

Return ONLY a valid JSON object with this structure:
{{
    "questions": [
        {{
            "type": "multiple_choice",
            "question": "What is...",
            "options": {{
                "a": "Option A",
                "b": "Option B",
                "c": "Option C",
                "d": "Option D"
            }},
            "correct_answer": "a",
            "explanation": "Brief explanation",
            "difficulty": "medium",
            "points": 1
        }},
        {{
            "type": "true_false",
            "question": "Statement here",
            "correct_answer": "true",
            "explanation": "Why this is true/false",
            "difficulty": "easy",
            "points": 1
        }},
        {{
            "type": "identification",
            "question": "What term describes...",
            "correct_answer": "The specific term",
            "explanation": "Explanation",
            "difficulty": "medium",
            "points": 1
        }}
    ]
}}

Ensure questions are:
- Directly based on the content provided
- Clear and unambiguous
- Appropriate difficulty level
- Academically sound
"""
        
        try:
            message = self.client.messages.create(
                model="claude-sonnet-4-20250514",
                max_tokens=4000,
                messages=[
                    {"role": "user", "content": prompt}
                ]
            )
            
            response_text = message.content[0].text
            
            # Clean up the response to extract JSON
            if "```json" in response_text:
                response_text = response_text.split("```json")[1].split("```")[0]
            elif "```" in response_text:
                response_text = response_text.split("```")[1].split("```")[0]
            
            return json.loads(response_text.strip())
        
        except Exception as e:
            raise Exception(f"AI extraction failed: {str(e)}")
    
    def process_questionnaire(self, questionnaire, question_types: List[str]):
        """Main method to process uploaded questionnaire"""
        from questionnaires.models import ExtractedQuestion, QuestionType
        
        # Extract text from file
        file_path = questionnaire.file.path
        content = self.extract_text_from_file(file_path)
        
        if not content.strip():
            raise Exception("No text content could be extracted from the file")
        
        # Use AI to extract questions
        extracted_data = self.extract_questions_with_ai(content, question_types)
        
        # Save extracted questions to database
        created_questions = []
        for q_data in extracted_data.get('questions', []):
            try:
                question_type = QuestionType.objects.get(name=q_data['type'])
                
                question = ExtractedQuestion.objects.create(
                    questionnaire=questionnaire,
                    question_type=question_type,
                    question_text=q_data['question'],
                    option_a=q_data.get('options', {}).get('a'),
                    option_b=q_data.get('options', {}).get('b'),
                    option_c=q_data.get('options', {}).get('c'),
                    option_d=q_data.get('options', {}).get('d'),
                    correct_answer=q_data['correct_answer'],
                    explanation=q_data.get('explanation', ''),
                    difficulty=q_data.get('difficulty', 'medium'),
                    points=q_data.get('points', 1)
                )
                created_questions.append(question)
            except QuestionType.DoesNotExist:
                continue
            except Exception as e:
                print(f"Error creating question: {str(e)}")
                continue
        
        return created_questions